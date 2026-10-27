import docx
import os
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

doc = docx.Document("Bao_Cao_ASL_Preprocessing_CNN.docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

headings_1 = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']
headings_2 = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2']
headings_3 = [p.text for p in doc.paragraphs if p.style.name == 'Heading 3']

print(f"\n--- HEADING 1 ({len(headings_1)}) ---")
for h in headings_1:
    print(f"  * {h}")

print(f"\n--- HEADING 2 ({len(headings_2)}) ---")
for h in headings_2:
    print(f"  - {h}")

# Check tables
print(f"\n--- TABLES ({len(doc.tables)}) ---")
for i, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns)
    first_cell = tbl.cell(0, 0).text.replace('\n', ' ')[:30]
    print(f"  Table {i+1}: {rows} rows x {cols} cols | First cell: '{first_cell}'")

# Check images (in inline shapes or document parts)
rels = doc.part.rels
images = [rel.target_ref for rel in rels.values() if "image" in rel.target_ref]
print(f"\nTotal Embedded Images: {len(images)}")
for img in images:
    print(f"  * {img}")
