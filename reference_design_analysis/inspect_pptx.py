import os
import sys
import json
import zipfile
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor

pptx_path = r"F:\OneDrive\hoc hanh\XLA\slide\Panorama_Stitching_14_Slides.pptx"

prs = Presentation(pptx_path)

slide_width = prs.slide_width
slide_height = prs.slide_height

print(f"Slide dimensions: {slide_width} x {slide_height} EMU")
print(f"Slide dimensions: {slide_width / 914400:.3f} x {slide_height / 914400:.3f} inches")
print(f"Slide dimensions: {slide_width / 12700:.2f} x {slide_height / 12700:.2f} pt")
aspect_ratio = slide_width / slide_height
print(f"Aspect ratio: {aspect_ratio:.3f} (16:9 is ~1.778)")

slides_data = []

def get_color_str(color_format):
    if color_format is None:
        return "None"
    try:
        if color_format.type == 1: # RGB
            return f"#{color_format.rgb}"
        elif color_format.type == 2: # Theme
            return f"Theme({color_format.theme_color})"
        else:
            return f"Type({color_format.type})"
    except Exception as e:
        return f"Error({e})"

for idx, slide in enumerate(prs.slides):
    s_info = {
        "slide_index": idx + 1,
        "notes": slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else "",
        "shapes": []
    }
    
    for s_idx, shape in enumerate(slide.shapes):
        shape_info = {
            "id": shape.shape_id,
            "name": shape.name,
            "type": str(shape.shape_type),
            "left_pt": round(shape.left / 12700, 2) if shape.left is not None else None,
            "top_pt": round(shape.top / 12700, 2) if shape.top is not None else None,
            "width_pt": round(shape.width / 12700, 2) if shape.width is not None else None,
            "height_pt": round(shape.height / 12700, 2) if shape.height is not None else None,
            "left_in": round(shape.left / 914400, 3) if shape.left is not None else None,
            "top_in": round(shape.top / 914400, 3) if shape.top is not None else None,
            "width_in": round(shape.width / 914400, 3) if shape.width is not None else None,
            "height_in": round(shape.height / 914400, 3) if shape.height is not None else None,
        }
        
        # Check fill
        try:
            fill = shape.fill
            if fill.type:
                shape_info["fill_type"] = str(fill.type)
                try:
                    shape_info["fill_color"] = get_color_str(fill.fore_color)
                except:
                    pass
        except:
            pass

        # Check line
        try:
            line = shape.line
            if line.fill.type:
                shape_info["line_color"] = get_color_str(line.color)
                shape_info["line_width_pt"] = round(line.width / 12700, 2) if line.width else None
        except:
            pass

        # Text frame
        if shape.has_text_frame:
            tf = shape.text_frame
            shape_info["text_frame"] = {
                "text": tf.text.strip(),
                "margin_left_pt": round(tf.margin_left / 12700, 2) if tf.margin_left else 0,
                "margin_right_pt": round(tf.margin_right / 12700, 2) if tf.margin_right else 0,
                "margin_top_pt": round(tf.margin_top / 12700, 2) if tf.margin_top else 0,
                "margin_bottom_pt": round(tf.margin_bottom / 12700, 2) if tf.margin_bottom else 0,
                "paragraphs": []
            }
            for p in tf.paragraphs:
                p_info = {
                    "text": p.text.strip(),
                    "alignment": str(p.alignment) if p.alignment else "None",
                    "level": p.level,
                    "runs": []
                }
                for r in p.runs:
                    r_info = {
                        "text": r.text,
                        "font_name": r.font.name,
                        "font_size_pt": r.font.size.pt if r.font.size else None,
                        "bold": r.font.bold,
                        "italic": r.font.italic,
                        "color": get_color_str(r.font.color) if r.font.color else None
                    }
                    p_info["runs"].append(r_info)
                shape_info["text_frame"]["paragraphs"].append(p_info)

        # Picture details
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                shape_info["image_filename"] = shape.image.filename
                shape_info["image_size"] = shape.image.size
                shape_info["image_content_type"] = shape.image.content_type
            except Exception as e:
                shape_info["image_error"] = str(e)
                
        # Group shapes
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            shape_info["group_shapes_count"] = len(shape.shapes)
            group_children = []
            for child in shape.shapes:
                child_data = {
                    "name": child.name,
                    "type": str(child.shape_type),
                    "left_pt": round(child.left / 12700, 2) if child.left is not None else None,
                    "top_pt": round(child.top / 12700, 2) if child.top is not None else None,
                    "width_pt": round(child.width / 12700, 2) if child.width is not None else None,
                    "height_pt": round(child.height / 12700, 2) if child.height is not None else None,
                    "text": child.text_frame.text.strip() if child.has_text_frame else ""
                }
                group_children.append(child_data)
            shape_info["children"] = group_children
            
        s_info["shapes"].append(shape_info)
        
    slides_data.append(s_info)

with open(r"reference_design_analysis\slides_data.json", "w", encoding="utf-8") as f:
    json.dump(slides_data, f, ensure_ascii=False, indent=2)

print(f"Extracted data for {len(prs.slides)} slides to slides_data.json.")
